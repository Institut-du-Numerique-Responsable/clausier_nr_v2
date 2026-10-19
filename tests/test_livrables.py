# -*- coding: utf-8 -*-
"""Tests des livrables : classeur Excel, site, cohérence des identifiants entre formats."""
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote, urlparse

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import build_excel  # noqa: E402
import build_site  # noqa: E402
import clausier  # noqa: E402

TMP = Path(tempfile.mkdtemp(prefix="clausier-tests-"))
CLAUSES = clausier.published(clausier.load_clauses())
IDS = [c.id for c in CLAUSES]

# budgets de poids (octets) : un dépassement doit être justifié
BUDGET = {"html": 60_000, "css": 15_000, "js": 15_000, "json": 200_000}


def setUpModule():
    build_excel.build(TMP / "clausier.xlsx", preselect=["COM-Q-13", "MAT-Q-02", "IA-Q-06"])
    build_site.EXCEL = TMP / "clausier.xlsx"
    build_site.Site(TMP / "site").build()


def tearDownModule():
    shutil.rmtree(TMP, ignore_errors=True)


class TestExcel(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(TMP / "clausier.xlsx")

    def test_feuilles(self):
        self.assertEqual(self.wb.sheetnames, ["Présentation", "Catalogue", "Sélection",
                                              "Clauses retenues", "Lexique", "Références", "Crédits"])

    def test_catalogue_complet_et_fidele(self):
        ws = self.wb["Catalogue"]
        head = [c.value for c in ws[1]]
        self.assertEqual([ws.cell(r, 1).value for r in range(2, ws.max_row + 1)], IDS)
        champ = {r["titre"]: r["champ"] for r in clausier.load_yaml("colonnes.yml")["rubriques"]}
        col = {k: head.index(h) + 1 for k, h in [("commentaire", "Commentaire"), ("exigences", "Exigences"),
               ("livrables", "Livrables"), ("questions", "Questions"),
               ("niveau_1", "Niveau 1 (Must have)"), ("niveau_2", "Niveau 2 (Nice to have)")]}
        for r, c in enumerate(CLAUSES, start=2):
            self.assertEqual(ws.cell(r, 2).value, c["titre"])
            for titre, md in c.rubriques:
                self.assertEqual(ws.cell(r, col[champ[titre]]).value, clausier.md_to_plain(md, c.notes),
                                 f"{c.id} {titre}")
        self.assertIn("Catalogue", ws.tables)

    def test_selection_liste_deroulante_et_formules(self):
        ws = self.wb["Sélection"]
        dvs = ws.data_validations.dataValidation
        self.assertEqual(len(dvs), 1)
        self.assertEqual(dvs[0].formula1, '"Oui,Non"')
        self.assertEqual(str(dvs[0].sqref), f"F4:F{len(IDS) + 3}")
        self.assertTrue(str(ws["G4"].value).startswith("=IF(F4=\"Oui\""))
        ret = self.wb["Clauses retenues"]
        self.assertIn("MATCH(1,", ret["B3"].value)
        for row in ret.iter_rows(min_row=3):
            for cell in row:
                v = str(cell.value or "")
                self.assertNotRegex(v, r"\b(FILTER|XLOOKUP|SORT|UNIQUE|SEQUENCE)\(")

    def test_credits_paul_saulnier(self):
        textes = " ".join(str(c.value) for row in self.wb["Crédits"].iter_rows() for c in row if c.value)
        self.assertIn("Paul Saulnier — Contribution initiale à la structuration et à la préparation "
                      "du classeur Excel du Clausier Numérique Responsable V2", textes)
        self.assertIn("CC BY 4.0", textes)
        self.assertIn(clausier.load_yaml("metadata.yml")["licence_contenu"]["attribution"], textes)

    def test_reproductible(self):
        autre = TMP / "bis.xlsx"
        build_excel.build(autre, preselect=["COM-Q-13", "MAT-Q-02", "IA-Q-06"])
        self.assertEqual(autre.read_bytes(), (TMP / "clausier.xlsx").read_bytes())

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice absent")
    def test_recalcul_libreoffice(self):
        """Les clauses pré-cochées remontent dans « Clauses retenues », sans erreur."""
        out = TMP / "lo"
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out),
                        str(TMP / "clausier.xlsx")], check=True, capture_output=True, timeout=180)
        wb = openpyxl.load_workbook(out / "clausier.xlsx", data_only=True)
        self.assertEqual(wb["Sélection"]["C2"].value, 3)
        ret = wb["Clauses retenues"]
        self.assertEqual([ret.cell(r, 2).value for r in range(3, 6)], ["COM-Q-13", "MAT-Q-02", "IA-Q-06"])
        self.assertIn(ret.cell(6, 2).value, (None, ""))
        self.assertEqual(ret["C5"].value, "Sobriété du Modèle")
        erreurs = [c.coordinate for ws in wb for row in ws.iter_rows() for c in row
                   if isinstance(c.value, str) and c.value.startswith("#")]
        self.assertEqual(erreurs, [])


class Liens(HTMLParser):
    def __init__(self):
        super().__init__()
        self.refs, self.h1, self.lang, self.ext, self.buttons = [], 0, None, [], []
        self._btn = None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "html":
            self.lang = a.get("lang")
        if tag == "h1":
            self.h1 += 1
        for k in ("href", "src"):
            if a.get(k):
                self.refs.append(a[k])
        if tag in ("script", "link") and (a.get("src") or a.get("href") or "").startswith("http"):
            self.ext.append(a.get("src") or a.get("href"))
        if tag == "button":
            self._btn = ""

    def handle_data(self, data):
        if self._btn is not None:
            self._btn += data

    def handle_endtag(self, tag):
        if tag == "button" and self._btn is not None:
            self.buttons.append(self._btn.strip())
            self._btn = None


class TestSite(unittest.TestCase):
    site = TMP / "site"

    def pages(self):
        return sorted(self.site.rglob("*.html"))

    def test_une_page_par_clause(self):
        for cid in IDS:
            self.assertTrue((self.site / f"clausier/clauses/{cid}/index.html").exists(), cid)

    def test_liens_internes_et_ressources(self):
        cassés = []
        for p in self.pages():
            parser = Liens()
            parser.feed(p.read_text(encoding="utf-8"))
            for ref in parser.refs:
                u = urlparse(ref)
                if u.scheme or ref.startswith(("#", "mailto:")):
                    continue
                cible = (p.parent / unquote(u.path)).resolve()
                if cible.is_dir():
                    cible = cible / "index.html"
                if not cible.exists():
                    cassés.append(f"{p.relative_to(self.site)} -> {ref}")
        self.assertEqual(cassés, [])

    def test_accessibilite_de_base(self):
        for p in self.pages():
            parser = Liens()
            parser.feed(p.read_text(encoding="utf-8"))
            rel = p.relative_to(self.site)
            self.assertEqual(parser.lang, "fr", rel)
            self.assertEqual(parser.h1, 1, rel)
            self.assertTrue(all(parser.buttons), f"{rel} : bouton sans texte")

    def test_aucune_ressource_externe(self):
        for p in self.pages():
            parser = Liens()
            parser.feed(p.read_text(encoding="utf-8"))
            self.assertEqual(parser.ext, [], p)

    def test_compatible_csp_sans_inline(self):
        """L'hébergeur applique script-src 'self' et style-src 'self' : rien d'inline."""
        for p in self.pages():
            h = p.read_text(encoding="utf-8")
            self.assertNotRegex(h, r"<script(?![^>]*\bsrc=)[^>]*>", p)
            self.assertNotIn("<style", h, p)
            self.assertNotRegex(h, r"\sstyle=\"", p)
            self.assertNotRegex(h, r"\son[a-z]+=\"", p)
        self.assertTrue((self.site / ".htaccess").exists())

    def test_budgets_de_poids(self):
        for p in self.site.rglob("*"):
            ext = p.suffix.lstrip(".")
            if p.is_file() and ext in BUDGET:
                self.assertLessEqual(p.stat().st_size, BUDGET[ext], f"{p.name} dépasse {BUDGET[ext]} octets")

    def test_donnees_json_coherentes(self):
        d = json.loads((self.site / "donnees/clauses.json").read_text(encoding="utf-8"))
        self.assertEqual([c["id"] for c in d["clauses"]], IDS)
        par_id = {c.id: c for c in CLAUSES}
        for c in d["clauses"]:
            self.assertEqual(c["rubriques"], [list(r) for r in par_id[c["id"]].rubriques])

    def test_meme_identifiant_meme_clause(self):
        """Markdown, site et Excel désignent la même clause par le même identifiant."""
        ws = openpyxl.load_workbook(TMP / "clausier.xlsx")["Catalogue"]
        excel = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(2, ws.max_row + 1)}
        for c in CLAUSES:
            page = (self.site / f"clausier/clauses/{c.id}/index.html").read_text(encoding="utf-8")
            h1 = re.search(r"<h1>(.*?)</h1>", page).group(1)
            self.assertEqual(build_site.html.unescape(h1), c["titre"])
            self.assertEqual(excel[c.id], c["titre"])

    def test_credits_site(self):
        page = (self.site / "credits/index.html").read_text(encoding="utf-8")
        self.assertIn('id="reutiliser"', page)
        self.assertIn("Institut du Numérique Responsable (INR) et French Tech Corporate Community", page)
        self.assertIn("<b>Paul Saulnier</b> — Contribution initiale à la structuration et à la "
                      "préparation du classeur\nExcel du Clausier Numérique Responsable V2", page)


if __name__ == "__main__":
    unittest.main()
